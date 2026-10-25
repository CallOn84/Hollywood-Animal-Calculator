#!/usr/bin/env python3
"""Build GameVariables.db for the Hollywood Animal Quick Mode calculator (v3.0.0 WIP).

Reads the game's JSON configs (TagData, TagCompatibilityData, GenrePairs), takes display
names / ordering from Quick Mode Prototype.xlsx (Tag List + Tag Values sheets) and writes
the SQLite schema that Code/3. Persistance/EagerLoading.js and the Quick Mode business
files query:
    meta(key, value)                                   -> data_version
    tags(tag_id, tag_key, name, category, sort_order, art_value, commercial_value)
    pair_scores(tag_id, partner_id, score)
    genre_pairs(genre_id, partner_id, art_value, commercial_value)
Usage: build_game_variables_db.py <Data/Configs dir> <Quick Mode Prototype.xlsx> <output.db>
"""
import hashlib, json, os, re, sqlite3, sys
import openpyxl

CATEGORIES = ["Genre", "Setting", "Protagonist", "Antagonist",
              "Supporting Character", "Theme", "Event", "Finale"]
TAG_LIST_COLUMNS = {"Genre": "Genre", "Setting": "Setting", "Protagonist": "Protagonist",
                    "Antagonist": "Antagonist", "Supporting Character": "Supporting Characters",
                    "Finale": "Finale"}   # Theme + Event share the "Themes and Events" column
WB_CATEGORY = {"Supporting Characters": "Supporting Character"}  # sheet label -> code label


def load_json(path):
    with open(path, encoding="utf-8-sig") as f:
        return json.load(f)


def json_category(key, tag):
    cid = tag.get("CategoryID")
    if cid == "SupportingCharacter":
        return "Supporting Character"
    if cid == "Theme" and tag.get("category") == 32:      # 32 = event flag, 16 = theme flag
        return "Event"
    if cid:
        return cid
    return "Genre" if tag["type"] == 0 else "Setting"


def norm(name):
    name = re.sub(r"<[^>]+>", "", name)                     # strip <nobr> markup
    name = re.sub(r"\s*\([^)]*\)\s*$", "", name)            # strip "(Genre)" style suffix
    return re.sub(r"\s+", " ", name).strip().lower()


def main(cfg_dir, xlsx_path, out_path):
    tag_json = load_json(os.path.join(cfg_dir, "TagData.json"))
    compat_json = load_json(os.path.join(cfg_dir, "TagCompatibilityData.json"))
    genre_json = load_json(os.path.join(cfg_dir, "GenrePairs.json"))

    # ---- workbook: names, categories, order ------------------------------------------------
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    wb_tags = {}   # (category, norm name) -> display name
    for row in wb["Tag Values"].iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        cat = WB_CATEGORY.get(row[1], row[1])
        k = (cat, norm(row[0]))
        assert k not in wb_tags, f"duplicate workbook tag {k}"
        wb_tags[k] = row[0]

    # tag list order -> sort_order (Theme/Event keep their relative order in the shared column)
    tl = wb["Tag List"]
    headers = [c.value for c in tl[1]]
    order = {}   # (category, display name) -> 1-based position
    for ci, header in enumerate(headers, start=1):
        names = [tl.cell(row=r, column=ci).value for r in range(2, tl.max_row + 1)]
        names = [n for n in names if n]
        if header == "Themes and Events":
            for cat in ("Theme", "Event"):
                members = [n for n in names if (cat, norm(n)) in wb_tags and wb_tags[(cat, norm(n))] == n]
                for i, n in enumerate(members, start=1):
                    order[(cat, n)] = i
        else:
            cat = next(c for c, h in TAG_LIST_COLUMNS.items() if h == header)
            for i, n in enumerate(names, start=1):
                order[(cat, n)] = i

    # ---- game localisation: JSON key -> English name (used to match to workbook) -----------
    loc_path = os.path.normpath(os.path.join(cfg_dir, "..", "Localization", "ENG", "NON_EVENT.json"))
    loc = load_json(loc_path)
    loc_name = lambda key: loc["locStrings"][loc["IdMap"][key]]

    # ---- build tag rows -----------------------------------------------------------------
    rows = []
    used = set()
    for key, tag in tag_json.items():
        cat = json_category(key, tag)
        wb_key = (cat, norm(loc_name(key)))
        if wb_key not in wb_tags:
            raise SystemExit(f"No workbook tag for {key} ({cat}, {loc_name(key)!r})")
        used.add(wb_key)
        name = wb_tags[wb_key]
        rows.append(dict(tag_key=key, name=name, category=cat,
                         sort_order=order[(cat, name)],
                         art_value=float(tag["artValue"]),
                         commercial_value=float(tag["commercialValue"])))
    assert used == set(wb_tags), "workbook tags with no JSON counterpart"
    rows.sort(key=lambda r: (CATEGORIES.index(r["category"]), r["sort_order"]))
    for i, r in enumerate(rows, start=1):
        r["tag_id"] = i
    id_of = {r["tag_key"]: r["tag_id"] for r in rows}

    # ---- data_version: fingerprint of the three source files ---------------------------------
    h = hashlib.sha256()
    for fn in ("TagData.json", "TagCompatibilityData.json", "GenrePairs.json"):
        with open(os.path.join(cfg_dir, fn), "rb") as f:
            h.update(f.read())
    data_version = "3.0.0-" + h.hexdigest()[:12]

    # ---- write database -----------------------------------------------------------------
    if os.path.exists(out_path):
        os.remove(out_path)
    con = sqlite3.connect(out_path)
    con.executescript("""
    CREATE TABLE meta (
        key   TEXT PRIMARY KEY NOT NULL,
        value TEXT NOT NULL
    );
    CREATE TABLE tags (
        tag_id           INTEGER PRIMARY KEY NOT NULL,
        tag_key          TEXT NOT NULL UNIQUE,
        name             TEXT NOT NULL,
        category         TEXT NOT NULL,
        sort_order       INTEGER NOT NULL,
        art_value        REAL NOT NULL,
        commercial_value REAL NOT NULL,
        UNIQUE (category, sort_order)
    );
    CREATE TABLE pair_scores (
        tag_id     INTEGER NOT NULL REFERENCES tags(tag_id),
        partner_id INTEGER NOT NULL REFERENCES tags(tag_id),
        score      REAL NOT NULL,
        PRIMARY KEY (tag_id, partner_id)
    ) WITHOUT ROWID;
    CREATE TABLE genre_pairs (
        genre_id         INTEGER NOT NULL REFERENCES tags(tag_id),
        partner_id       INTEGER NOT NULL REFERENCES tags(tag_id),
        art_value        REAL NOT NULL,
        commercial_value REAL NOT NULL,
        PRIMARY KEY (genre_id, partner_id)
    ) WITHOUT ROWID;
    """)
    con.executemany("INSERT INTO meta VALUES (?, ?)", [("data_version", data_version)])
    con.executemany(
        "INSERT INTO tags VALUES (:tag_id, :tag_key, :name, :category, :sort_order, :art_value, :commercial_value)", rows)
    con.executemany("INSERT INTO pair_scores VALUES (?, ?, ?)",
                    [(id_of[a], id_of[b], float(s)) for a, v in compat_json.items() for b, s in v.items()])
    # GenrePairs.json: Item1 = commercial value, Item2 = art value
    con.executemany("INSERT INTO genre_pairs VALUES (?, ?, ?, ?)",
                    [(id_of[g], id_of[p], float(v["Item2"]), float(v["Item1"]))
                     for g, d in genre_json.items() for p, v in d.items()])
    con.commit()
    con.execute("VACUUM")
    con.close()
    print(f"wrote {out_path}  data_version={data_version}")


if __name__ == "__main__":
    if len(sys.argv) != 4:
        sys.exit(__doc__)
    main(*sys.argv[1:])
